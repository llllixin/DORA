"""经营数据文件入库：解析 + 校验（引擎输入规范格式）。

格式：CSV/XLSX，表头 metric_key,label,dimension,value,unit。
支持指标 key：margin/returns/orders/revenue/aov/east_orders/new_sku/high_value/supplier_price。
"""
import csv
import io
import zipfile
from itertools import islice
from xml.etree.ElementTree import ParseError
from typing import Any

ALLOWED_KEYS = {
    "margin", "returns", "orders", "revenue", "aov",
    "east_orders", "new_sku", "high_value", "supplier_price",
}
REQUIRED_COLS = ["metric_key", "label", "value"]
OPTIONAL_COLS = ["dimension", "unit"]
MAX_DATASET_ROWS = 10000
MAX_XLSX_EXPANDED_BYTES = 16 * 1024 * 1024
MAX_COLUMNS = 256


class IngestValidationError(Exception):
    """文件格式/内容校验失败（路由映射 400）。"""


class IngestCapacityError(IngestValidationError):
    status_code = 413


def _bounded_rows(rows, max_rows, truncate):
    result=[]
    for row in islice(rows,max_rows if truncate else max_rows+1):
        if len(result)>=max_rows:
            raise IngestCapacityError('数据超过下游 10000 行处理上限，请拆分后导入')
        result.append(dict(row))
    return result


def _rows_from_csv(content: bytes, max_rows, truncate) -> list[dict[str, str]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise IngestValidationError("文件为空或无表头")
    if len(reader.fieldnames)>MAX_COLUMNS:raise IngestCapacityError('文件列数超过 256 列上限')
    return _bounded_rows(reader,max_rows,truncate)


def _rows_from_xlsx(content: bytes, max_rows, truncate) -> list[dict[str, str]]:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist())>MAX_XLSX_EXPANDED_BYTES:
            raise IngestCapacityError('XLSX 解压后超过 16MiB 上限，请拆分后导入')
    from openpyxl import load_workbook  # 延迟导入，仅上传 xlsx 时加载
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        raise IngestValidationError("xlsx 没有工作表")
    try:
        if ws.max_column and ws.max_column>MAX_COLUMNS:raise IngestCapacityError('文件列数超过 256 列上限')
        it=ws.iter_rows(values_only=True,max_col=MAX_COLUMNS,max_row=max_rows+2)
        try:header=[str(h).strip() if h is not None else '' for h in next(it)]
        except StopIteration:raise IngestValidationError('xlsx 为空或无表头')
        def values():
            for row in it:
                yield {header[i]:('' if v is None else str(v).strip()) for i,v in enumerate(row) if header[i]}
        return [r for r in _bounded_rows(values(),max_rows,truncate) if any(r.values())]
    finally:wb.close()


def parse_dataset(filename: str, content: bytes) -> list[dict[str, Any]]:
    """解析并校验；返回 series 行（可直接入库）。失败抛 IngestValidationError。"""
    raw = read_rows(content, filename)

    if not raw:
        raise IngestValidationError("文件内容为空")
    items: list[dict[str, Any]] = []
    for i, row in enumerate(raw, start=2):
        for col in REQUIRED_COLS:
            if col not in row or row[col] is None or str(row[col]).strip() == "":
                raise IngestValidationError(f"第 {i} 行缺少必填列 {col}")
        key = row["metric_key"].strip()
        if key not in ALLOWED_KEYS:
            raise IngestValidationError(f"第 {i} 行指标 key 不支持：{key}")
        try:
            value = float(row["value"])
        except ValueError:
            raise IngestValidationError(f"第 {i} 行 value 非数值：{row['value']}")
        if not (0 <= abs(value) <= 1e12):
            raise IngestValidationError(f"第 {i} 行 value 超出合理范围：{value}")
        items.append({
            "metric_key": key,
            "label": str(row["label"]).strip(),
            "dimension": str(row.get("dimension") or "").strip(),
            "value": value,
            "unit": str(row.get("unit") or "").strip(),
        })
    return _validate_items(items)


def read_rows(content: bytes, filename: str, *, max_rows=MAX_DATASET_ROWS, truncate=False) -> list[dict[str, str]]:
    """仅解析（不做值校验），供 preview/映射使用。"""
    name = filename.lower()
    try:
        if name.endswith(".csv"):
            return _rows_from_csv(content,max_rows,truncate)
        if name.endswith(".xlsx"):
            return _rows_from_xlsx(content,max_rows,truncate)
    except (UnicodeError, csv.Error, zipfile.BadZipFile, KeyError, ValueError, OSError, ParseError) as exc:
        raise IngestValidationError("文件无法解析，请检查 UTF-8 CSV 或有效 XLSX 格式") from exc
    raise IngestValidationError("仅支持 .csv / .xlsx")


def preview_rows(content: bytes, filename: str, max_rows: int = 10) -> dict[str, Any]:
    raw = read_rows(content, filename,max_rows=max_rows,truncate=True)
    if not raw:
        raise IngestValidationError("文件内容为空或无有效数据行")
    columns = list(dict.fromkeys(k for row in raw for k in row))
    return {"columns": columns, "sampleRows": raw[:max_rows]}


def items_from_mapping(raw: list[dict[str, str]], mapping: dict[str, Any]) -> list[dict[str, Any]]:
    """按映射把任意列布局行转成 series 项（校验语义与规范直传一致）。"""
    for name in ("metric_key", "label_column", "value_column", "dimension_column", "unit"):
        if mapping.get(name) is not None and not isinstance(mapping[name], str):
            raise IngestValidationError(f"映射字段 {name} 必须是文本列名")
    metric = (mapping.get("metric_key") or "").strip()
    if metric not in ALLOWED_KEYS:
        raise IngestValidationError(f"指标 key 不支持：{metric or '(空)'}")
    label_col = (mapping.get("label_column") or "").strip()
    value_col = (mapping.get("value_column") or "").strip()
    dim_col = (mapping.get("dimension_column") or "").strip() or None
    unit_col = (mapping.get("unit") or "").strip() or None
    if not label_col or not value_col:
        raise IngestValidationError("必须指定日期列和数值列")
    if not raw:
        raise IngestValidationError("文件内容为空")

    first = raw[0]
    for col in [label_col, value_col, dim_col, unit_col]:
        if col and col not in first:
            raise IngestValidationError(f"文件缺少列：{col}")

    items: list[dict[str, Any]] = []
    for i, row in enumerate(raw, start=2):
        if all(not str(row.get(c) or "").strip() for c in row):
            continue
        label = str(row.get(label_col) or "").strip()
        if not label:
            raise IngestValidationError(f"第 {i} 行 label 为空")
        try:
            value = float(str(row.get(value_col) or "").replace(",", ""))
        except (TypeError, ValueError):
            raise IngestValidationError(f"第 {i} 行 {value_col} 非数值：{row.get(value_col)}")
        if not (0 <= abs(value) <= 1e12):
            raise IngestValidationError(f"第 {i} 行 value 超出合理范围：{value}")
        items.append({
            "metric_key": metric,
            "label": label,
            "dimension": str(row.get(dim_col) or "").strip() if dim_col else "",
            "value": value,
            "unit": str(row.get(unit_col) or "").strip() if unit_col else "",
        })
    if not items:
        raise IngestValidationError("没有可入库的数据行")
    return _validate_items(items)


DEFAULT_DIMENSIONS = {
    "margin": "全国", "orders": "全国", "revenue": "全国", "aov": "全国",
    "high_value": "全国", "east_orders": "华东", "new_sku": "全区域",
}


def _validate_items(items: list[dict]) -> list[dict]:
    """直传与映射共用输入契约：先验证整批，再允许写库。"""
    seen = set()
    for row in items:
        key = row["metric_key"]
        if not row["dimension"]:
            if key not in DEFAULT_DIMENSIONS:
                raise IngestValidationError(f"{key} 必须提供门店或供应商维度")
            row["dimension"] = DEFAULT_DIMENSIONS[key]
        if len(row["label"]) > 32 or len(row["dimension"]) > 32 or len(row["unit"]) > 16:
            raise IngestValidationError("日期/维度最长 32 字符，单位最长 16 字符")
        identity = (key, row["label"], row["dimension"])
        if identity in seen:
            raise IngestValidationError(f"重复记录：{key} / {row['label']} / {row['dimension']}")
        seen.add(identity)
    return items
